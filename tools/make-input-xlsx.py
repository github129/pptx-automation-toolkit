"""templates/スライド入力.xlsx を作り直すスクリプト (配布物の更新用)。

  python3 tools/make-input-xlsx.py

「スライド」シートには samples/all-types.json と同じ内容 (全種類の見本) を入れる。
"""
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

ROOT = Path(__file__).resolve().parent.parent
deck = json.loads((ROOT / "samples" / "all-types.json").read_text(encoding="utf-8"))

DESIGNS = ["ビジネス", "コーポレート", "ナチュラル", "モダン", "ポップ", "ミニマル", "エレガント", "テック"]
TYPES = ["表紙", "目次", "章", "箇条書き", "カード", "数字", "ステップ", "タイムライン", "比較", "表",
         "棒グラフ", "積み上げ棒グラフ", "横棒グラフ", "折れ線グラフ", "面グラフ", "円グラフ", "ドーナツグラフ",
         "画像", "引用", "メッセージ", "終わり"]
WRAP = Alignment(wrap_text=True, vertical="top")
HEAD = Font(bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor="0F2A4A")

wb = Workbook()

# ---------- 使い方 ----------
ws = wb.active
ws.title = "使い方"
guide = [
    ("PowerPoint 自動作成 — 入力シート", ""),
    ("", ""),
    ("手順", "1. このファイルを OneDrive の /PowerPoint/ に保存する\n2. 「設定」シートでデザインなどを選ぶ\n3. 「スライド」シートに 1 行 = 1 スライドで書く\n4. Power Automate のフロー「PowerPoint作成」を実行する"),
    ("", ""),
    ("種類", "「内容」の書き方"),
    ("表紙", "サブタイトル (改行で 2 行まで)。空なら作成者と日付が入る"),
    ("目次", "空なら「章」のタイトルから自動で作る。書く場合は 1 行 = 1 項目"),
    ("章", "サブタイトル。章番号 01・02… は自動で付く"),
    ("箇条書き", "1 行 = 1 項目。行頭の全角スペースで 2 段目。**太字** も使える。「画像」を入れると右に画像"),
    ("カード", "見出し | 説明 を 1 行ずつ (2～8 個)"),
    ("数字", "数値 | 説明 | 補足 を 1 行ずつ (最大 6 個)"),
    ("ステップ", "ステップ名 | 説明 を 1 行ずつ (最大 6 個)"),
    ("タイムライン", "時期 | 出来事 を 1 行ずつ (最大 7 個)"),
    ("比較", "見出し行の下に「・」付きで項目。見出しが 2 つなら左右、3 つなら 3 列"),
    ("表", "列を | で区切る。1 行目が見出し。11 行を超えると次のスライドに分かれる"),
    ("棒グラフ / 横棒グラフ / 折れ線グラフ / 面グラフ", "項目 | 値 を 1 行ずつ。1 行目に系列名を書くと複数系列。「---」の下の行はグラフ右の箇条書き"),
    ("積み上げ棒グラフ", "1 行目に 項目名 | 系列 1 | 系列 2 …、2 行目以降に値"),
    ("円グラフ / ドーナツグラフ", "項目 | 値 を 1 行ずつ。割合 (%) は自動で表示"),
    ("画像", "「画像」列に OneDrive のパス (例 /PowerPoint/images/image.png)。「内容」を書くと左に箇条書き"),
    ("引用", "発言を書き、最後に「— 名前」の行を書くと発言者になる"),
    ("メッセージ", "伝えたい一言の補足。タイトルが中央に大きく出る"),
    ("終わり", "連絡先など"),
    ("", ""),
    ("ヒント", "・項目は | (全角の ｜ も可) で区切る\n・セル内の改行は Alt + Enter\n・色を変えたいときは「設定」の「メインカラー」に #E11D48 のように入力\n・詳しくは docs/DeckSpec.md を参照"),
]
for r in guide:
    ws.append(list(r))
ws["A1"].font = Font(bold=True, size=16)
for c in ("A3", "A5", "B5", "A24"):
    ws[c].font = Font(bold=True)
for row in ws.iter_rows(min_row=3):
    for c in row:
        c.alignment = WRAP
ws.column_dimensions["A"].width = 34
ws.column_dimensions["B"].width = 90

# ---------- 設定 ----------
ws = wb.create_sheet("設定")
settings = [
    ("デザイン", deck["デザイン"], "ビジネス / コーポレート / ナチュラル / モダン / ポップ / ミニマル / エレガント / テック"),
    ("テンプレート", "", "社内テンプレート (.pptx / .potx) の OneDrive パス。空なら「デザイン」を使う"),
    ("メインカラー", "", "例 #2563EB。空ならデザインの色"),
    ("サブカラー", "", "例 #0EA5A4。空ならデザインの色"),
    ("フォント", "", "例 メイリオ。空ならデザインのフォント"),
    ("ロゴ", "", "全スライド右上に入れる画像の OneDrive パス"),
    ("フッター", deck["フッター"], "全スライド下部の文字"),
    ("ページ番号", deck["ページ番号"], "はい / いいえ"),
    ("作成者", deck["作成者"], "表紙のサブタイトルとファイルのプロパティに使う"),
    ("日付", deck["日付"], "表紙に入る日付"),
    ("切り替え", deck["切り替え"], "フェード / プッシュ / ワイプ / カバー / なし"),
    ("ファイル名", deck["ファイル名"], "{timestamp} は実行日時に置き換わる。空なら表紙のタイトル"),
    ("保存先", deck["保存先"], "PowerPoint を保存する OneDrive のフォルダー"),
]
ws.append(["項目", "値", "説明"])
for r in settings:
    ws.append(list(r))
for row in ws.iter_rows(min_row=2):
    row[1].number_format = "@"  # 日付などを文字列のまま保持する
ws.add_table(Table(displayName="設定", ref=f"A1:C{len(settings) + 1}",
                   tableStyleInfo=TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)))
ws.column_dimensions["A"].width = 16
ws.column_dimensions["B"].width = 40
ws.column_dimensions["C"].width = 80
dv = DataValidation(type="list", formula1='"' + ",".join(DESIGNS) + '"', allow_blank=True)
ws.add_data_validation(dv)
dv.add("B2")
dv = DataValidation(type="list", formula1='"はい,いいえ"', allow_blank=True)
ws.add_data_validation(dv)
dv.add("B8")
dv = DataValidation(type="list", formula1='"フェード,プッシュ,ワイプ,カバー,なし"', allow_blank=True)
ws.add_data_validation(dv)
dv.add("B11")

# ---------- スライド ----------
ws = wb.create_sheet("スライド")
cols = ["種類", "タイトル", "内容", "画像", "ノート"]
ws.append(cols)
for s in deck["slides"]:
    ws.append([s.get(c, "") for c in cols])
last = len(deck["slides"]) + 1
for row in ws.iter_rows(min_row=2):
    for c in row:
        c.alignment = WRAP
ws.add_table(Table(displayName="スライド", ref=f"A1:E{last}",
                   tableStyleInfo=TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)))
for col, w in zip("ABCDE", (16, 34, 70, 30, 50)):
    ws.column_dimensions[col].width = w
ws.freeze_panes = "A2"
dv = DataValidation(type="list", formula1="=種類一覧!$A$1:$A$%d" % len(TYPES), allow_blank=True)
ws.add_data_validation(dv)
dv.add("A2:A500")

# 種類のドロップダウン用 (非表示)
ws = wb.create_sheet("種類一覧")
for t in TYPES:
    ws.append([t])
ws.sheet_state = "hidden"

wb.active = wb.sheetnames.index("スライド")
for w in wb.worksheets:
    w.sheet_view.tabSelected = w.title == "スライド"
out = ROOT / "templates" / "スライド入力.xlsx"
wb.save(out)
print(out)
